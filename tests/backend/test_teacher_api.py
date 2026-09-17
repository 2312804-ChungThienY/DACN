"""API Giảng viên 2.1, 2.2 – phân quyền theo môn, bảng cảnh báo, Modal, can thiệp, tiến độ, so sánh, export."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.backend.conftest import API, TEST_MARK

RANK = {"RED": 0, "YELLOW": 1, "GREEN": 2}


@pytest.fixture()
def course(accounts):
    return accounts["course"]


def base(course) -> str:
    return f"{API}/teacher/courses/{course['course_id']}"


def test_teacher_sees_only_own_courses(client, teacher_headers, accounts):
    items = client.get(f"{API}/teacher/courses", headers=teacher_headers).json()["items"]
    assert sorted(i["course_id"] for i in items) == sorted(c["course_id"] for c in accounts["teacher_courses"])
    by_term = client.get(f"{API}/teacher/courses", params={"academic_year": "2026-2027", "semester": 1},
                         headers=teacher_headers).json()["items"]
    assert len(by_term) == len(items)
    assert client.get(f"{API}/teacher/courses", params={"academic_year": "2025-2026", "semester": 2},
                      headers=teacher_headers).json()["items"] == []


def test_course_permission(client, teacher_headers, other_teacher_headers, admin_headers, course):
    for suffix in ("/risk-summary", "/students", "/progress", "/export"):
        assert client.get(base(course) + suffix, headers=other_teacher_headers).status_code == 403
        assert client.get(base(course) + suffix, headers=admin_headers).status_code == 403
    assert client.get(f"{API}/teacher/courses/999999/students", headers=teacher_headers).status_code == 404


def test_risk_summary(client, teacher_headers, course, snapshot):
    body = client.get(base(course) + "/risk-summary", headers=teacher_headers).json()
    assert body["snapshot_date"] == snapshot.isoformat()
    assert body["red"] == course["red_count"] and body["yellow"] == course["yellow_count"]
    assert body["red"] + body["yellow"] + body["green"] == body["total"] == course["enrolled"]
    for key, count in body["by_criterion"].items():
        assert body["by_criterion_red"][key] <= count <= body["total"]


def test_students_default_order_and_totals(client, teacher_headers, course):
    body = client.get(base(course) + "/students", params={"page_size": 200}, headers=teacher_headers).json()
    assert body["total"] == course["enrolled"]
    keys = [(RANK[s["risk_level"]], -s["risk_score"]) for s in body["items"]]
    assert keys == sorted(keys)
    first = body["items"][0]
    assert first["risk_level"] == "RED" and first["risk_reasons"]
    assert first["missing_assignments"] == f"{first['missing_assignments_count']}/{first['total_assignments']}"


def test_students_filters(client, teacher_headers, course):
    url = base(course) + "/students"
    red = client.get(url, params={"risk_level": "RED", "page_size": 200}, headers=teacher_headers).json()
    assert red["total"] == course["red_count"] and {s["risk_level"] for s in red["items"]} == {"RED"}
    red_yellow = client.get(url, params=[("risk_level", "RED"), ("risk_level", "YELLOW")], headers=teacher_headers).json()
    assert red_yellow["total"] == course["red_count"] + course["yellow_count"]

    absent = client.get(url, params={"criterion": "attendance", "page_size": 200}, headers=teacher_headers).json()
    assert absent["total"] > 0 and all(s["score_attendance"] > 0 for s in absent["items"])

    name = red["items"][0]["student_name"]
    found = client.get(url, params={"search": name.split()[-1].lower()}, headers=teacher_headers).json()
    assert any(s["student_name"] == name for s in found["items"])

    never_first = client.get(url, params={"sort": "days_since_last_access", "order": "desc", "page_size": 200},
                             headers=teacher_headers).json()["items"]
    days = [s["days_since_last_access"] for s in never_first]
    nulls = days.count(None)
    assert days[:nulls] == [None] * nulls
    assert days[nulls:] == sorted(days[nulls:], reverse=True)
    assert client.get(url, params={"sort": "email"}, headers=teacher_headers).status_code == 422
    assert client.get(url, params={"risk_level": "BLUE"}, headers=teacher_headers).status_code == 422


def test_student_detail(client, teacher_headers, course, db, snapshot):
    student = client.get(base(course) + "/students", params={"risk_level": "RED"}, headers=teacher_headers).json()["items"][0]
    body = client.get(f"{base(course)}/students/{student['user_id']}", headers=teacher_headers).json()
    assert body["risk_level"] == student["risk_level"] and body["snapshot_date"] == snapshot.isoformat()
    assert body["weekly_activity"] and body["weekly_activity"][0]["week"] == 1
    assert body["assignment_detail"] and {"name", "status", "late"} <= set(body["assignment_detail"][0])
    history_days = db.execute("""SELECT COUNT(*) FROM analytics.fact_student_risk_summary
                                 WHERE course_id = %s AND user_id = %s""", (course["course_id"], student["user_id"])).fetchone()[0]
    assert len(body["risk_history"]) == history_days
    assert body["risk_history"][-1]["date"] == snapshot.isoformat()
    assert body["class_avg_access"] is not None
    assert client.get(f"{base(course)}/students/999999999", headers=teacher_headers).status_code == 404


def test_create_intervention(client, teacher_headers, other_teacher_headers, course):
    student = client.get(base(course) + "/students", headers=teacher_headers).json()["items"][0]
    url = f"{base(course)}/students/{student['user_id']}/interventions"
    res = client.post(url, json={"action_type": "NOTE", "content": f"{TEST_MARK} Đã nhắc trên lớp  "},
                      headers=teacher_headers)
    assert res.status_code == 201, res.text
    created = res.json()
    detail = client.get(f"{base(course)}/students/{student['user_id']}", headers=teacher_headers).json()
    mine = next(i for i in detail["interventions"] if i["id"] == created["id"])
    assert mine["action_type"] == "NOTE" and mine["content"] == f"{TEST_MARK} Đã nhắc trên lớp"
    assert mine["risk_level_at_time"] == student["risk_level"]

    assert client.post(url, json={"action_type": "SMS", "content": TEST_MARK}, headers=teacher_headers).status_code == 422
    assert client.post(url, json={"action_type": "NOTE", "content": "x" * 2001}, headers=teacher_headers).status_code == 422
    assert client.post(url, json={"action_type": "NOTE", "content": TEST_MARK}, headers=other_teacher_headers).status_code == 403
    missing = f"{base(course)}/students/999999999/interventions"
    assert client.post(missing, json={"action_type": "NOTE", "content": TEST_MARK}, headers=teacher_headers).status_code == 404


def test_progress_daily_and_weekly(client, teacher_headers, course, db, snapshot):
    days = db.execute("SELECT COUNT(*) FROM analytics.fact_course_progress_daily WHERE course_id = %s",
                      (course["course_id"],)).fetchone()[0]
    daily = client.get(base(course) + "/progress", headers=teacher_headers).json()
    assert len(daily["series"]) == days
    assert daily["summary"]["stat_date"] == snapshot.isoformat() == daily["series"][-1]["period_end"]
    assert daily["summary"]["submission_rate"] == pytest.approx(float(course["assignment_submission_rate"]))
    assert daily["sections"] and all(0 <= s["completion_rate"] <= 100 for s in daily["sections"])

    weekly = client.get(base(course) + "/progress", params={"granularity": "week"}, headers=teacher_headers).json()
    periods = [p["period"] for p in weekly["series"]]
    assert len(periods) == len(set(periods)) == 9                     # 07/09 – 06/11/2026 = 9 tuần ISO
    by_day = {p["period_end"]: p for p in daily["series"]}
    for week in weekly["series"]:
        assert week["active_students_7d"] == by_day[week["period_end"]]["active_students_7d"]
    assert sum(w["total_views_day"] for w in weekly["series"]) == sum(d["total_views_day"] for d in daily["series"])
    assert client.get(base(course) + "/progress", params={"granularity": "month"}, headers=teacher_headers).status_code == 422


def test_compare(client, teacher_headers, course):
    student = client.get(base(course) + "/students", params={"risk_level": "RED"}, headers=teacher_headers).json()["items"][0]
    body = client.get(f"{base(course)}/students/{student['user_id']}/compare", headers=teacher_headers).json()
    metrics = {m["key"]: m for m in body["metrics"]}
    assert set(metrics) == {"access_14d", "total_access", "submission_rate", "attendance_rate", "completion_rate"}
    assert metrics["access_14d"]["student"] == student["access_count_14d"]
    assert metrics["submission_rate"]["class_avg"] == pytest.approx(float(course["assignment_submission_rate"]))
    assert metrics["attendance_rate"]["class_avg"] == pytest.approx(float(course["attendance_rate"]))
    if student["total_assignments"]:
        expected = 100 * student["submitted_assignments"] / student["total_assignments"]
        assert metrics["submission_rate"]["student"] == pytest.approx(expected, abs=0.01)


def test_teacher_export(client, teacher_headers, course):
    res = client.get(base(course) + "/export", headers=teacher_headers)
    assert res.status_code == 200
    assert "attachment" in res.headers["content-disposition"] and course["course_shortname"] in res.headers["content-disposition"]
    sheet = load_workbook(BytesIO(res.content), read_only=True)["Báo cáo"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Khoa" and len(rows) - 1 == course["enrolled"]

    red = client.get(base(course) + "/export", params={"risk_level": "RED"}, headers=teacher_headers)
    red_rows = list(load_workbook(BytesIO(red.content), read_only=True)["Báo cáo"].iter_rows(values_only=True))[1:]
    assert len(red_rows) == course["red_count"] and {r[5] for r in red_rows} == {"Nguy cơ cao"}

    pdf = client.get(base(course) + "/export", params={"format": "pdf"}, headers=teacher_headers)
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")

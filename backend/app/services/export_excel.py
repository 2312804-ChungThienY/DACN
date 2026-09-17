"""Xuất báo cáo Excel bằng openpyxl ở chế độ write-only (ít bộ nhớ, nhanh với hàng chục nghìn dòng)."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.services.formatting import RISK_COLORS, ExportDocument, excel_value

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NUMBER_FORMATS = {"int": "#,##0", "days": "#,##0", "percent": "0.00", "decimal": "0.00",
                  "date": "DD/MM/YYYY", "datetime": "DD/MM/YYYY HH:MM"}
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
WRAP = Alignment(wrap_text=True, vertical="top")


def build_xlsx(doc: ExportDocument) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Báo cáo")
    ws.freeze_panes = "A2"
    for index, column in enumerate(doc.columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = max(8.0, round(column.width * 16, 1))
    last = get_column_letter(len(doc.columns))
    ws.auto_filter.ref = f"A1:{last}{len(doc.rows) + 1}"

    header = []
    for column in doc.columns:
        cell = WriteOnlyCell(ws, value=column.title)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, Alignment(wrap_text=True, vertical="center")
        header.append(cell)
    ws.append(header)

    risk_fonts = {level: Font(bold=True, color=color) for level, color in RISK_COLORS.items()}
    for row in doc.rows:
        cells = []
        for column in doc.columns:
            raw = row.get(column.key)
            cell = WriteOnlyCell(ws, value=excel_value(raw, column.kind))
            if column.kind in NUMBER_FORMATS and raw is not None:
                cell.number_format = NUMBER_FORMATS[column.kind]
            if column.kind == "risk_level" and raw in risk_fonts:
                cell.font = risk_fonts[raw]
            if column.kind == "list":
                cell.alignment = WRAP
            cells.append(cell)
        ws.append(cells)

    info = wb.create_sheet("Thông tin")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 60
    title = WriteOnlyCell(info, value=doc.title)
    title.font = Font(bold=True, size=14)
    info.append([title])
    info.append(["Thời điểm xuất", doc.generated_at.strftime("%d/%m/%Y %H:%M")])
    info.append(["Số dòng", len(doc.rows)])
    for label, value in doc.filters:
        info.append([label, value])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
